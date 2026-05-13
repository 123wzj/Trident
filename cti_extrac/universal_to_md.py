import hashlib
import os
import re
from pathlib import Path
from typing import Optional, Union
from urllib.parse import urlparse

import html2text
import requests
from docx import Document

from agent.env_utils import MINERU_API_KEY


class UniversalToMarkdown:
    """Convert URLs and common document formats to clean Markdown."""

    def __init__(self, use_mineru_for_web: bool = True, default_save_dir: str = "./cti"):
        self.use_mineru_for_web = use_mineru_for_web
        self.default_save_dir = default_save_dir

    def convert(self, source: Union[str, Path], save_to_dir: Optional[str] = None) -> str:
        source = str(source)
        if self._is_url(source):
            markdown = self._from_url(source)
            source_name = self._url_to_filename(source, markdown)
        elif os.path.isfile(source):
            markdown = self._from_file(source)
            source_name = Path(source).stem
        else:
            raise ValueError(f"Invalid input: {source} is neither a valid URL nor a file path")

        target_dir = self.default_save_dir if save_to_dir is None else save_to_dir
        if target_dir:
            self._save_markdown(markdown, source_name, target_dir)
        return markdown

    def _url_to_filename(self, url: str, content: Optional[str] = None) -> str:
        if content:
            for line in content.strip().split("\n")[:20]:
                line = line.strip()
                if line.startswith("# "):
                    return self._sanitize_filename(line.lstrip("#").strip())
                if line and not line.startswith(("*", "-", ">", "```", "|", "#", "[", "!")):
                    return self._sanitize_filename(line[:80])

        parsed = urlparse(url)
        path_parts = [part for part in parsed.path.split("/") if part]
        if path_parts:
            return self._sanitize_filename(path_parts[-1])
        return f"{parsed.netloc}_{hashlib.md5(url.encode()).hexdigest()[:8]}"

    def _sanitize_filename(self, title: str, max_length: int = 100) -> str:
        title = re.sub(r'[<>:"/\\|?*\n\r]', "", title)
        title = re.sub(r"\s+", "_", title.strip())
        title = re.sub(r"[^\w\-_.]", "", title)
        return title[:max_length].rstrip("_-.") or "untitled"

    def _save_markdown(self, content: str, filename: str, save_dir: str):
        save_path = Path(save_dir)
        save_path.mkdir(parents=True, exist_ok=True)
        output_file = save_path / f"{filename}.md"
        output_file.write_text(content, encoding="utf-8")
        print(f"Saved Markdown to: {output_file.absolute()}")

    def _is_url(self, value: str) -> bool:
        try:
            parsed = urlparse(value)
            return bool(parsed.scheme and parsed.netloc)
        except Exception:
            return False

    def _from_url(self, url: str) -> str:
        if self.use_mineru_for_web and MINERU_API_KEY and not url.endswith("/"):
            try:
                print("Extracting with MinerU API...")
                markdown = self._mineru_api_convert(url, "markdown")
                print(f"MinerU extraction succeeded: {len(markdown)} characters")
                return self._post_clean(markdown)
            except Exception as exc:
                print(f"MinerU API failed: {exc}")
                print("Falling back to html2text...")

        try:
            markdown = self._html2text_extract(url)
            print(f"html2text extraction succeeded: {len(markdown)} characters")
            return self._post_clean(markdown)
        except Exception as exc:
            raise RuntimeError(f"Failed to process URL {url}: {exc}") from exc

    def _html2text_extract(self, url: str) -> str:
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
            )
        }
        response = requests.get(url, headers=headers, timeout=30)
        response.raise_for_status()
        return self._html_converter(ignore_images=True).handle(response.text)

    def _from_file(self, filepath: str) -> str:
        ext = Path(filepath).suffix.lower()
        converters = {
            ".pdf": self._from_pdf,
            ".docx": self._from_docx,
            ".html": self._from_html,
            ".txt": self._from_txt,
        }
        if ext not in converters:
            raise ValueError(f"Unsupported file format: {ext}")
        return converters[ext](filepath)

    def _from_pdf(self, filepath: str) -> str:
        if not MINERU_API_KEY:
            raise RuntimeError("MinerU API key not configured. Cannot process PDF.")

        url = self._upload_to_transfer_sh(filepath) if not self._is_url(filepath) else filepath
        print("Processing PDF with MinerU API...")
        markdown = self._mineru_api_convert(url, "markdown")
        print(f"PDF conversion succeeded: {len(markdown)} characters")
        return self._post_clean(markdown)

    def _mineru_api_convert(self, source: str, output_format: str = "markdown") -> str:
        payload = {"url": source, "output_format": output_format, "model_version": "vlm"}
        headers = {"Authorization": f"Bearer {MINERU_API_KEY}", "Content-Type": "application/json"}
        response = requests.post("https://api.mineru.net/v1/extract", json=payload, headers=headers, timeout=60)

        if response.status_code != 200:
            raise RuntimeError(f"API request failed (HTTP {response.status_code}): {response.text}")
        result = response.json()
        if result.get("status") != "success":
            raise RuntimeError(f"MinerU API error: {result.get('message')}")
        return result["data"]["content"]

    def _upload_to_transfer_sh(self, filepath: str) -> str:
        print("Uploading PDF to temporary storage...")
        with open(filepath, "rb") as file:
            response = requests.post("https://transfer.sh/", files={"file": file}, timeout=60)
        if response.status_code != 200:
            raise RuntimeError(f"Failed to upload to transfer.sh: {response.text}")
        url = response.text.strip()
        print(f"Upload succeeded: {url}")
        return url

    def _from_docx(self, filepath: str) -> str:
        try:
            doc = Document(filepath)
            blocks = []
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                if paragraph.style.name.startswith("Heading"):
                    level = min(int(paragraph.style.name[-1]), 6) if paragraph.style.name[-1].isdigit() else 2
                    blocks.append(f"{'#' * level} {text}")
                else:
                    blocks.append(text)

            blocks.extend(filter(None, (self._table_to_md(table) for table in doc.tables)))
            return self._post_clean("\n\n".join(blocks))
        except Exception as exc:
            raise RuntimeError(f"Failed to parse DOCX {filepath}: {exc}") from exc

    def _from_html(self, filepath: str) -> str:
        with open(filepath, "r", encoding="utf-8") as file:
            return self._post_clean(self._html_converter(ignore_images=False).handle(file.read()))

    def _from_txt(self, filepath: str) -> str:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as file:
            paragraphs = [part.strip() for part in file.read().split("\n\n") if part.strip()]
        return "\n\n".join(paragraphs)

    def _table_to_md(self, table) -> str:
        rows = []
        try:
            for index, row in enumerate(table.rows):
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                rows.append("| " + " | ".join(cells) + " |")
                if index == 0:
                    rows.append("| " + " | ".join(["---"] * len(cells)) + " |")
        except Exception:
            return ""
        return "\n".join(rows)

    def _post_clean(self, markdown: str, highlight_ioc: bool = False) -> str:
        if not markdown:
            return ""
        markdown = re.sub(r"\n{3,}", "\n\n", markdown)
        markdown = re.sub(r"<!--.*?-->", "", markdown, flags=re.DOTALL)
        markdown = re.sub(r"Page \d+ of \d+", "", markdown, flags=re.IGNORECASE)
        return self._highlight_ioc_safe(markdown).strip() if highlight_ioc else markdown.strip()

    def _highlight_ioc_safe(self, markdown: str) -> str:
        lines = []
        in_code_block = False
        for line in markdown.split("\n"):
            if line.strip().startswith("```"):
                in_code_block = not in_code_block
                lines.append(line)
                continue
            if not in_code_block and "`" not in line:
                line = re.sub(
                    r"(?<!\d)(\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3})(?!\d)",
                    lambda match: f"`IP:{match.group(1)}`" if self._is_valid_ip(match.group(1)) else match.group(1),
                    line,
                )
                line = re.sub(r"\b([a-fA-F0-9]{64})\b", r"`HASH:\1`", line)
            lines.append(line)
        return "\n".join(lines)

    def _is_valid_ip(self, ip: str) -> bool:
        try:
            return all(0 <= int(part) <= 255 for part in ip.split("."))
        except ValueError:
            return False

    @staticmethod
    def _html_converter(ignore_images: bool) -> html2text.HTML2Text:
        converter = html2text.HTML2Text()
        converter.ignore_links = False
        converter.ignore_images = ignore_images
        converter.ignore_emphasis = False
        converter.ignore_tables = False
        converter.body_width = 0
        converter.skip_internal_links = False
        converter.inline_links = True
        converter.wrap_links = False
        converter.default_image_alt = "image"
        return converter


if __name__ == "__main__":
    converter = UniversalToMarkdown(use_mineru_for_web=True, default_save_dir="./cti")
    url = ""
    converter.convert(url)
